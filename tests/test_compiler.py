import io
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.core.exceptions import InvalidColorError, ImageProcessingError
from app.engine.compiler import compiler
from app.engine.coordinates import get_column_bounds, get_row_bounds
from app.engine.images import (
    decode_and_validate_base64_image,
    calculate_contain_bounds,
)
from app.engine.theme import parse_hex_color, calculate_adaptive_font_size
from app.schemas.deck import (
    DeckPayload,
    DeckTheme,
    TitleSlide,
    CardsSlide,
    CardItem,
    TableSlide,
    TableData,
    SplitSlide,
    ColumnContent,
    ImageSlide,
    ProcessSlide,
    ProcessStep,
    ChartSlide,
    ChartData,
    ChartSeries,
)

# 1x1 pixel valid base64 PNG
TINY_PNG_BASE64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
ZIP_MAGIC_SIGNATURE = b"PK\x03\x04"


class TestCoordinates:
    def test_column_bounds_single_col(self):
        lefts, width = get_column_bounds(total_cols=1, left=1.0, total_width=11.333, gap=0.35)
        assert len(lefts) == 1
        assert lefts[0] == 1.0
        assert width == pytest.approx(11.333)

    def test_column_bounds_multiple_cols(self):
        lefts, width = get_column_bounds(total_cols=3, left=1.0, total_width=11.333, gap=0.35)
        assert len(lefts) == 3
        # col_w = (11.333 - (0.35 * 2)) / 3 = 10.633 / 3 = 3.54433...
        expected_w = (11.333 - 0.70) / 3
        assert width == pytest.approx(expected_w)
        assert lefts[0] == pytest.approx(1.0)
        assert lefts[1] == pytest.approx(1.0 + expected_w + 0.35)
        assert lefts[2] == pytest.approx(1.0 + (expected_w + 0.35) * 2)
        # Right edge
        assert (lefts[2] + width) == pytest.approx(1.0 + 11.333)

    def test_row_bounds(self):
        tops, height = get_row_bounds(total_rows=3, top=2.0, total_height=4.0, gap=0.5)
        assert len(tops) == 3
        expected_h = (4.0 - 1.0) / 3
        assert height == pytest.approx(expected_h)
        assert tops[0] == pytest.approx(2.0)


class TestThemeAndColors:
    def test_parse_hex_color_valid(self):
        color1 = parse_hex_color("#3B82F6")
        assert (color1[0], color1[1], color1[2]) == (0x3B, 0x82, 0xF6)

        color2 = parse_hex_color("0F172A")
        assert (color2[0], color2[1], color2[2]) == (0x0F, 0x17, 0x2A)

    def test_parse_hex_color_shorthand(self):
        color = parse_hex_color("#FFF")
        assert (color[0], color[1], color[2]) == (255, 255, 255)

    def test_parse_hex_color_invalid(self):
        with pytest.raises(InvalidColorError):
            parse_hex_color("not-a-color")

        with pytest.raises(InvalidColorError):
            parse_hex_color("#12345")

    def test_adaptive_font_sizing(self):
        # Short text retains base font size
        assert calculate_adaptive_font_size("Short text", base_pt=12, min_pt=9, threshold_chars=110) == Pt(12)
        # Medium text stepped down
        med_text = "x" * 150
        assert calculate_adaptive_font_size(med_text, base_pt=12, min_pt=9, threshold_chars=110, step_down_rate=35) == Pt(10)
        # Very long text clamped at min_pt
        long_text = "x" * 300
        assert calculate_adaptive_font_size(long_text, base_pt=12, min_pt=9, threshold_chars=110, step_down_rate=35) == Pt(9)


class TestImageEngine:
    def test_decode_and_validate_png(self):
        stream, w, h = decode_and_validate_base64_image(TINY_PNG_BASE64)
        assert isinstance(stream, io.BytesIO)
        assert w == 1.0
        assert h == 1.0

    def test_decode_with_data_uri_prefix(self):
        data_uri = f"data:image/png;base64,{TINY_PNG_BASE64}"
        stream, w, h = decode_and_validate_base64_image(data_uri)
        assert w == 1.0
        assert h == 1.0

    def test_decode_corrupted_base64(self):
        with pytest.raises(ImageProcessingError):
            decode_and_validate_base64_image("invalid_base64_!@#$")

    def test_contain_bounds_aspect_ratio(self):
        # Image is 200x100 (2:1). Box is 4x4.
        fit_l, fit_t, fit_w, fit_h = calculate_contain_bounds(
            orig_w=200, orig_h=100, box_left=1.0, box_top=1.0, box_w=4.0, box_h=4.0
        )
        assert fit_w == pytest.approx(4.0)
        assert fit_h == pytest.approx(2.0)
        assert fit_l == pytest.approx(1.0)
        assert fit_t == pytest.approx(2.0)  # centered vertically: 1.0 + (4 - 2)/2


class TestLayoutCompilation:
    """Every slide layout must verify binary output is non-empty and starts with PK\x03\x04."""

    def test_compile_title_slide(self):
        payload = DeckPayload(
            title="Title Deck",
            slides=[
                TitleSlide(
                    kicker="SERIES A",
                    title="Autonomous Systems",
                    subtitle="Compiler Architecture",
                )
            ],
        )
        buffer = compiler.compile(payload)
        data = buffer.getvalue()
        assert len(data) > 0
        assert data.startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_cards_slide(self):
        for num_cards in [1, 2, 3, 4]:
            cards = [
                CardItem(
                    header=f"Card {i}",
                    body="Card body narrative description.",
                    metric=f"{i * 10}x",
                    tag="METRIC",
                )
                for i in range(num_cards)
            ]
            payload = DeckPayload(
                title=f"Cards Deck {num_cards}",
                slides=[
                    CardsSlide(
                        kicker="HIGHLIGHTS",
                        title="Key Performance Indicators",
                        cards=cards,
                    )
                ],
            )
            buffer = compiler.compile(payload)
            data = buffer.getvalue()
            assert len(data) > 0
            assert data.startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_table_slide(self):
        payload = DeckPayload(
            title="Table Deck",
            slides=[
                TableSlide(
                    kicker="BENCHMARKS",
                    title="Performance Matrix",
                    table=TableData(
                        headers=["Engine", "Latency", "Memory"],
                        rows=[
                            ["DeckBridge", "< 200ms", "Zero Disk"],
                            ["Legacy Script", "> 3000ms", "Local File /tmp"],
                        ],
                    ),
                )
            ],
        )
        buffer = compiler.compile(payload)
        data = buffer.getvalue()
        assert len(data) > 0
        assert data.startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_split_slide(self):
        payload = DeckPayload(
            title="Split Deck",
            slides=[
                SplitSlide(
                    kicker="COMPARISON",
                    title="Before vs After",
                    left_column=ColumnContent(
                        header="Legacy Architecture",
                        body="Client-side DOM rendering.",
                        items=["Brittle iframe downloads", "Font clipping"],
                    ),
                    right_column=ColumnContent(
                        header="DeckBridge Engine",
                        body="Headless serverless compiler.",
                        items=["Sub-second in-memory stream", "100% OpenXML compliance"],
                    ),
                )
            ],
        )
        buffer = compiler.compile(payload)
        data = buffer.getvalue()
        assert len(data) > 0
        assert data.startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_image_slide(self):
        payload = DeckPayload(
            title="Image Deck",
            slides=[
                ImageSlide(
                    kicker="VISUALIZATION",
                    title="System Architecture Diagram",
                    image_base64=TINY_PNG_BASE64,
                    caption="Figure 1: In-memory presentation compilation pipeline.",
                )
            ],
        )
        buffer = compiler.compile(payload)
        data = buffer.getvalue()
        assert len(data) > 0
        assert data.startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_process_slides(self):
        steps = [
            ProcessStep(title="Ingest", body="Parse JSON schema", shape_type="rounded_rectangle"),
            ProcessStep(title="Validate", body="Pydantic v2 contract", shape_type="diamond"),
            ProcessStep(title="Render", body="Native OpenXML vector shapes", shape_type="rectangle"),
            ProcessStep(title="Stream", body="Stream io.BytesIO directly", shape_type="oval"),
        ]

        # Horizontal mode
        payload_h = DeckPayload(
            title="Horizontal Process",
            slides=[ProcessSlide(title="Sequential Pipeline", direction="horizontal", steps=steps)],
        )
        buf_h = compiler.compile(payload_h)
        assert buf_h.getvalue().startswith(ZIP_MAGIC_SIGNATURE)

        # Vertical mode
        payload_v = DeckPayload(
            title="Vertical Process",
            slides=[ProcessSlide(title="Vertical Stack", direction="vertical", steps=steps)],
        )
        buf_v = compiler.compile(payload_v)
        assert buf_v.getvalue().startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_chart_column_clustered(self):
        payload = DeckPayload(
            title="Column Chart Deck",
            slides=[
                ChartSlide(
                    kicker="PERFORMANCE",
                    title="Quarterly Growth Trajectory",
                    chart=ChartData(
                        chart_type="column_clustered",
                        categories=["Q1", "Q2", "Q3", "Q4"],
                        series=[
                            ChartSeries(name="2025", values=[1.2, 1.8, 2.4, 3.1]),
                            ChartSeries(name="2026", values=[2.0, 3.2, 4.5, 6.0]),
                        ],
                        show_legend=True,
                        show_data_labels=True,
                    ),
                )
            ],
        )
        buf = compiler.compile(payload)
        data = buf.getvalue()
        assert len(data) > 0
        assert data.startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_chart_bar_and_line(self):
        for c_type in ["bar_clustered", "line"]:
            payload = DeckPayload(
                title=f"{c_type} Chart Deck",
                slides=[
                    ChartSlide(
                        title=f"{c_type} Test",
                        chart=ChartData(
                            chart_type=c_type,
                            categories=["Jan", "Feb", "Mar"],
                            series=[ChartSeries(name="Series A", values=[10.0, 20.0, 30.0])],
                        ),
                    )
                ],
            )
            buf = compiler.compile(payload)
            assert buf.getvalue().startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_chart_pie(self):
        payload = DeckPayload(
            title="Pie Chart Deck",
            slides=[
                ChartSlide(
                    title="Market Share Breakdown",
                    chart=ChartData(
                        chart_type="pie",
                        categories=["North America", "Europe", "APAC"],
                        series=[ChartSeries(name="Share", values=[45.0, 35.0, 20.0])],
                    ),
                )
            ],
        )
        buf = compiler.compile(payload)
        assert buf.getvalue().startswith(ZIP_MAGIC_SIGNATURE)

    def test_compile_chart_custom_series_colors(self):
        payload = DeckPayload(
            title="Custom Colors Chart",
            slides=[
                ChartSlide(
                    title="Themed Chart",
                    chart=ChartData(
                        chart_type="column_clustered",
                        categories=["A", "B"],
                        series=[
                            ChartSeries(name="Revenue", values=[10.0, 20.0], color="10B981"),
                            ChartSeries(name="Cost", values=[5.0, 8.0], color="EF4444"),
                        ],
                    ),
                )
            ],
        )
        buf = compiler.compile(payload)
        assert buf.getvalue().startswith(ZIP_MAGIC_SIGNATURE)


class TestOpenXmlCompliance:
    def test_presentation_dimensions_16_9(self):
        payload = DeckPayload(
            title="Dimension Test",
            slides=[TitleSlide(title="Dimension Test Slide")],
        )
        buffer = compiler.compile(payload)
        prs = Presentation(buffer)
        assert prs.slide_width == Inches(13.333)
        assert prs.slide_height == Inches(7.5)

    def test_zero_margins_and_word_wrap(self):
        payload = DeckPayload(
            title="Margin Test",
            slides=[
                CardsSlide(
                    title="Zero Margins Test",
                    cards=[CardItem(header="Card 1", body="Multi-line body narrative text")],
                )
            ],
        )
        buffer = compiler.compile(payload)
        prs = Presentation(buffer)
        slide = prs.slides[0]

        # Scan text frames and verify zero margins
        found_tf = False
        for shape in slide.shapes:
            if shape.has_text_frame:
                found_tf = True
                tf = shape.text_frame
                assert tf.margin_top == Inches(0)
                assert tf.margin_bottom == Inches(0)
                assert tf.margin_left == Inches(0)
                assert tf.margin_right == Inches(0)
                assert tf.word_wrap is True
        assert found_tf is True
