from typing import Optional
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.engine.coordinates import (
    HEADER_LEFT,
    HEADER_TOP,
    HEADER_WIDTH,
    HEADER_HEIGHT,
)
from app.engine.theme import ResolvedTheme, apply_zero_margins_and_wrap
from app.schemas.deck import BaseSlide


class BaseLayoutBuilder:
    """Base class for all slide layout builders."""

    def __init__(self, prs, theme: ResolvedTheme):
        self.prs = prs
        self.theme = theme

    def create_blank_slide(self, slide_model: BaseSlide):
        """Instantiate a slide using blank slide layout (index 6) and configure background."""
        # Use blank slide layout
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)

        # Resolve background and text contrast colors
        bg_color, text_color, muted_color = self.theme.resolve_slide_colors(
            slide_model.background_override
        )

        try:
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = bg_color
        except Exception:
            # Fallback: draw full-bleed background shape if slide.background fill is restricted
            from pptx.enum.shapes import MSO_SHAPE
            bg_rect = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), self.prs.slide_width, self.prs.slide_height
            )
            bg_rect.fill.solid()
            bg_rect.fill.fore_color.rgb = bg_color
            bg_rect.line.fill.background()  # No border

        return slide, bg_color, text_color, muted_color

    def render_header(
        self,
        slide,
        kicker: Optional[str],
        title: str,
        subtitle: Optional[str],
        text_color: RGBColor,
        muted_color: RGBColor,
    ) -> None:
        """Render standardized top header canvas with kicker, title, and subtitle."""
        textbox = slide.shapes.add_textbox(
            Inches(HEADER_LEFT),
            Inches(HEADER_TOP),
            Inches(HEADER_WIDTH),
            Inches(HEADER_HEIGHT),
        )
        tf = textbox.text_frame
        apply_zero_margins_and_wrap(tf)

        first_p = True

        # 1. Kicker
        if kicker:
            p_kicker = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_kicker.text = kicker.upper()
            p_kicker.font.name = self.theme.font_family
            p_kicker.font.size = Pt(11)
            p_kicker.font.bold = True
            p_kicker.font.color.rgb = self.theme.accent
            p_kicker.space_after = Pt(4)
            first_p = False

        # 2. Title
        p_title = tf.paragraphs[0] if first_p else tf.add_paragraph()
        p_title.text = title
        p_title.font.name = self.theme.font_family
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = text_color
        first_p = False

        # 3. Subtitle
        if subtitle:
            p_title.space_after = Pt(4)
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = self.theme.font_family
            p_sub.font.size = Pt(14)
            p_sub.font.color.rgb = muted_color
