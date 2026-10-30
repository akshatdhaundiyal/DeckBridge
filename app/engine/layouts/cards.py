from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from app.engine.coordinates import (
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
    get_column_bounds,
)
from app.engine.images import decode_and_validate_base64_image, calculate_contain_bounds
from app.engine.layouts.base import BaseLayoutBuilder
from app.engine.theme import apply_zero_margins_and_wrap, parse_hex_color, calculate_adaptive_font_size
from app.schemas.deck import CardsSlide


class CardsLayoutBuilder(BaseLayoutBuilder):
    """Layout builder for Cards grid slides (1 to 4 cards)."""

    def render(self, slide_model: CardsSlide):
        slide, _, text_color, muted_color = self.create_blank_slide(slide_model)

        self.render_header(
            slide=slide,
            kicker=slide_model.kicker,
            title=slide_model.title,
            subtitle=slide_model.subtitle,
            text_color=text_color,
            muted_color=muted_color,
        )

        cards = slide_model.cards
        col_lefts, col_w = get_column_bounds(
            total_cols=len(cards),
            left=CONTENT_LEFT,
            total_width=CONTENT_WIDTH,
            gap=0.35,
        )

        card_top = CONTENT_TOP
        card_height = CONTENT_HEIGHT

        for idx, (card, col_left) in enumerate(zip(cards, col_lefts)):
            # 1. Card Container Shape
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(col_left),
                Inches(card_top),
                Inches(col_w),
                Inches(card_height),
            )
            if shape.has_text_frame:
                apply_zero_margins_and_wrap(shape.text_frame)
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.theme.card_bg
            shape.line.color.rgb = self.theme.card_border
            shape.line.width = Pt(1)

            # Determine accent color for this card
            card_accent = (
                parse_hex_color(card.accent_color) if card.accent_color else self.theme.accent
            )

            # 2. Card Content Textbox
            pad = 0.25
            inner_left = col_left + pad
            inner_top = card_top + pad
            inner_w = col_w - (pad * 2)

            # Check if there is an image
            has_image = bool(card.image_base64)
            text_height = (card_height * 0.55) if has_image else (card_height - (pad * 2))

            textbox = slide.shapes.add_textbox(
                Inches(inner_left),
                Inches(inner_top),
                Inches(inner_w),
                Inches(text_height),
            )
            tf = textbox.text_frame
            apply_zero_margins_and_wrap(tf)

            first_p = True

            # Category / Tag
            if card.tag:
                p_tag = tf.paragraphs[0] if first_p else tf.add_paragraph()
                p_tag.text = card.tag.upper()
                p_tag.font.name = self.theme.font_family
                p_tag.font.size = Pt(10)
                p_tag.font.bold = True
                p_tag.font.color.rgb = card_accent
                p_tag.space_after = Pt(4)
                first_p = False

            # Prominent Metric Callout (e.g. "$4.2B", "10x")
            if card.metric:
                p_metric = tf.paragraphs[0] if first_p else tf.add_paragraph()
                p_metric.text = card.metric
                p_metric.font.name = self.theme.font_family
                p_metric.font.size = Pt(32)
                p_metric.font.bold = True
                p_metric.font.color.rgb = card_accent
                p_metric.space_after = Pt(4)
                first_p = False

            # Headline / Header
            if card.display_header:
                p_hdr = tf.paragraphs[0] if first_p else tf.add_paragraph()
                p_hdr.text = card.display_header
                p_hdr.font.name = self.theme.font_family
                p_hdr.font.size = Pt(17)
                p_hdr.font.bold = True
                p_hdr.font.color.rgb = text_color
                p_hdr.space_after = Pt(6)
                first_p = False

            # Body narrative
            p_body = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_body.text = card.body
            p_body.font.name = self.theme.font_family
            p_body.font.size = calculate_adaptive_font_size(
                card.body, base_pt=12, min_pt=9, threshold_chars=110, step_down_rate=35
            )
            p_body.font.color.rgb = muted_color

            # 3. Optional Thumbnail Image
            if has_image:
                img_stream, orig_w, orig_h = decode_and_validate_base64_image(card.image_base64)
                img_box_top = inner_top + text_height + 0.1
                img_box_h = card_height - text_height - (pad * 2) - 0.1
                fit_left, fit_top, fit_w, fit_h = calculate_contain_bounds(
                    orig_w, orig_h, inner_left, img_box_top, inner_w, img_box_h
                )
                slide.shapes.add_picture(
                    img_stream,
                    Inches(fit_left),
                    Inches(fit_top),
                    Inches(fit_w),
                    Inches(fit_h),
                )

        return slide
