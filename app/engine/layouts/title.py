from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from app.engine.layouts.base import BaseLayoutBuilder
from app.engine.theme import apply_zero_margins_and_wrap
from app.schemas.deck import TitleSlide


class TitleLayoutBuilder(BaseLayoutBuilder):
    """Layout builder for Title / Hero slide."""

    def render(self, slide_model: TitleSlide):
        slide, _, text_color, muted_color = self.create_blank_slide(slide_model)

        # Center/Hero Layout Frame
        left = Inches(1.2)
        top = Inches(2.2)
        width = Inches(10.933)
        height = Inches(3.6)

        textbox = slide.shapes.add_textbox(left, top, width, height)
        tf = textbox.text_frame
        apply_zero_margins_and_wrap(tf)

        first_p = True

        # Kicker
        if slide_model.kicker:
            p_kicker = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_kicker.text = slide_model.kicker.upper()
            p_kicker.font.name = self.theme.font_family
            p_kicker.font.size = Pt(13)
            p_kicker.font.bold = True
            p_kicker.font.color.rgb = self.theme.accent
            p_kicker.space_after = Pt(12)
            first_p = False

        # Main Title
        p_title = tf.paragraphs[0] if first_p else tf.add_paragraph()
        p_title.text = slide_model.title
        p_title.font.name = self.theme.font_family
        p_title.font.size = Pt(42)
        p_title.font.bold = True
        p_title.font.color.rgb = text_color
        p_title.space_after = Pt(12)
        first_p = False

        # Subtitle
        if slide_model.subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = slide_model.subtitle
            p_sub.font.name = self.theme.font_family
            p_sub.font.size = Pt(18)
            p_sub.font.color.rgb = muted_color

        # Subtle brand accent line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left,
            Inches(1.8),
            Inches(0.8),
            Inches(0.06),
        )
        if line.has_text_frame:
            apply_zero_margins_and_wrap(line.text_frame)
        line.fill.solid()
        line.fill.fore_color.rgb = self.theme.accent
        line.line.fill.background()

        return slide
