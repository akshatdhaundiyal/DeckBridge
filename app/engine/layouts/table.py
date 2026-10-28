from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.engine.coordinates import (
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
)
from app.engine.layouts.base import BaseLayoutBuilder
from app.engine.theme import apply_zero_margins_and_wrap
from app.schemas.deck import TableSlide


class TableLayoutBuilder(BaseLayoutBuilder):
    """Layout builder for formatted OpenXML Table slides."""

    def render(self, slide_model: TableSlide):
        slide, _, text_color, muted_color = self.create_blank_slide(slide_model)

        self.render_header(
            slide=slide,
            kicker=slide_model.kicker,
            title=slide_model.title,
            subtitle=slide_model.subtitle,
            text_color=text_color,
            muted_color=muted_color,
        )

        table_data = slide_model.table
        headers = table_data.headers
        rows = table_data.rows

        num_rows = len(rows) + 1
        num_cols = len(headers)

        table_shape = slide.shapes.add_table(
            num_rows,
            num_cols,
            Inches(CONTENT_LEFT),
            Inches(CONTENT_TOP),
            Inches(CONTENT_WIDTH),
            Inches(CONTENT_HEIGHT),
        )
        table = table_shape.table

        # Configure column widths
        if table_data.column_widths and len(table_data.column_widths) == num_cols:
            for col_idx, col_w in enumerate(table_data.column_widths):
                table.columns[col_idx].width = Inches(col_w)
        else:
            default_w = CONTENT_WIDTH / num_cols
            for col_idx in range(num_cols):
                table.columns[col_idx].width = Inches(default_w)

        # Style Header Row
        for col_idx, header_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = header_text
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.theme.primary

            tf = cell.text_frame
            apply_zero_margins_and_wrap(tf)
            for p in tf.paragraphs:
                p.font.name = self.theme.font_family
                p.font.size = Pt(13)
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)

        # Style Data Rows with Zebra Striping
        white_rgb = RGBColor(255, 255, 255)
        for row_idx, row_data in enumerate(rows, start=1):
            is_zebra = (row_idx % 2 == 0)
            row_bg = self.theme.card_bg if is_zebra else white_rgb

            for col_idx in range(num_cols):
                cell = table.cell(row_idx, col_idx)
                cell_value = row_data[col_idx] if col_idx < len(row_data) else ""
                cell.text = str(cell_value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_bg

                tf = cell.text_frame
                apply_zero_margins_and_wrap(tf)
                for p in tf.paragraphs:
                    p.font.name = self.theme.font_family
                    p.font.size = Pt(12)
                    p.font.color.rgb = text_color

        return slide
