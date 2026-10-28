from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from app.engine.coordinates import (
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
    get_column_bounds,
)
from app.engine.images import decode_and_validate_base64_image, calculate_contain_bounds
from app.engine.layouts.base import BaseLayoutBuilder
from app.engine.theme import apply_zero_margins_and_wrap, parse_hex_color
from app.schemas.deck import SplitSlide, ColumnContent


class SplitLayoutBuilder(BaseLayoutBuilder):
    """Layout builder for 2-column comparative Split slides."""

    def render(self, slide_model: SplitSlide):
        slide, _, text_color, muted_color = self.create_blank_slide(slide_model)

        self.render_header(
            slide=slide,
            kicker=slide_model.kicker,
            title=slide_model.title,
            subtitle=slide_model.subtitle,
            text_color=text_color,
            muted_color=muted_color,
        )

        col_lefts, col_w = get_column_bounds(
            total_cols=2,
            left=CONTENT_LEFT,
            total_width=CONTENT_WIDTH,
            gap=0.5,
        )

        columns = [slide_model.left_column, slide_model.right_column]

        for col_left, col_content in zip(col_lefts, columns):
            self._render_column(
                slide=slide,
                col_content=col_content,
                left=col_left,
                top=CONTENT_TOP,
                width=col_w,
                height=CONTENT_HEIGHT,
                default_text_color=text_color,
                default_muted_color=muted_color,
            )

        return slide

    def _render_column(
        self,
        slide,
        col_content: ColumnContent,
        left: float,
        top: float,
        width: float,
        height: float,
        default_text_color,
        default_muted_color,
    ):
        # Container shape
        container = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        if container.has_text_frame:
            apply_zero_margins_and_wrap(container.text_frame)
        container.fill.solid()
        container.fill.fore_color.rgb = self.theme.card_bg
        container.line.color.rgb = self.theme.card_border
        container.line.width = Pt(1)

        pad = 0.3
        inner_left = left + pad
        inner_top = top + pad
        inner_w = width - (pad * 2)

        has_image = bool(col_content.image_base64)
        text_h = (height * 0.45) if has_image else (height - (pad * 2))

        col_accent = (
            parse_hex_color(col_content.accent_color)
            if col_content.accent_color
            else self.theme.accent
        )

        textbox = slide.shapes.add_textbox(
            Inches(inner_left),
            Inches(inner_top),
            Inches(inner_w),
            Inches(text_h),
        )
        tf = textbox.text_frame
        apply_zero_margins_and_wrap(tf)

        first_p = True

        # Column Header
        if col_content.header:
            p_hdr = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_hdr.text = col_content.header
            p_hdr.font.name = self.theme.font_family
            p_hdr.font.size = Pt(18)
            p_hdr.font.bold = True
            p_hdr.font.color.rgb = default_text_color
            p_hdr.space_after = Pt(8)
            first_p = False

        # Column Body
        if col_content.body:
            p_body = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_body.text = col_content.body
            p_body.font.name = self.theme.font_family
            p_body.font.size = Pt(13)
            p_body.font.color.rgb = default_muted_color
            p_body.space_after = Pt(8)
            first_p = False

        # Bullet Items
        if col_content.items:
            for item in col_content.items:
                p_item = tf.paragraphs[0] if first_p else tf.add_paragraph()
                p_item.text = f"•  {item}"
                p_item.font.name = self.theme.font_family
                p_item.font.size = Pt(12)
                p_item.font.color.rgb = default_text_color
                p_item.space_after = Pt(4)
                first_p = False

        # Embedded Image
        if has_image:
            img_stream, orig_w, orig_h = decode_and_validate_base64_image(col_content.image_base64)
            img_box_top = inner_top + text_h + 0.1
            img_box_h = height - text_h - (pad * 2) - 0.1
            fit_left, fit_top, fit_w, fit_h = calculate_contain_bounds(
                orig_w, orig_h, inner_left, img_box_top, inner_w, img_box_h
            )
            slide.shapes.add_picture(
                img_stream,
                Inches(fit_left),
                Inches(fit_top),
                Inches(fit_w),
                Inches(fit_h),
            )
