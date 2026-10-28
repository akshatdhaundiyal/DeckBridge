from pptx.util import Inches, Pt

from app.engine.coordinates import (
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
)
from app.engine.images import decode_and_validate_base64_image, calculate_contain_bounds
from app.engine.layouts.base import BaseLayoutBuilder
from app.engine.theme import apply_zero_margins_and_wrap
from app.schemas.deck import ImageSlide


class ImageLayoutBuilder(BaseLayoutBuilder):
    """Layout builder for dedicated showcase Image slides."""

    def render(self, slide_model: ImageSlide):
        slide, _, text_color, muted_color = self.create_blank_slide(slide_model)

        self.render_header(
            slide=slide,
            kicker=slide_model.kicker,
            title=slide_model.title,
            subtitle=slide_model.subtitle,
            text_color=text_color,
            muted_color=muted_color,
        )

        has_caption = bool(slide_model.caption)
        caption_height = 0.45 if has_caption else 0.0

        img_box_left = CONTENT_LEFT
        img_box_top = CONTENT_TOP
        img_box_w = CONTENT_WIDTH
        img_box_h = CONTENT_HEIGHT - caption_height

        # Decode base64 image and calculate aspect ratio containment bounds
        img_stream, orig_w, orig_h = decode_and_validate_base64_image(slide_model.image_base64)
        fit_left, fit_top, fit_w, fit_h = calculate_contain_bounds(
            orig_w, orig_h, img_box_left, img_box_top, img_box_w, img_box_h
        )

        slide.shapes.add_picture(
            img_stream,
            Inches(fit_left),
            Inches(fit_top),
            Inches(fit_w),
            Inches(fit_h),
        )

        # Render Caption
        if has_caption:
            caption_top = CONTENT_TOP + img_box_h + 0.1
            textbox = slide.shapes.add_textbox(
                Inches(CONTENT_LEFT),
                Inches(caption_top),
                Inches(CONTENT_WIDTH),
                Inches(caption_height),
            )
            tf = textbox.text_frame
            apply_zero_margins_and_wrap(tf)
            p = tf.paragraphs[0]
            p.text = slide_model.caption
            p.font.name = self.theme.font_family
            p.font.size = Pt(11)
            p.font.italic = True
            p.font.color.rgb = muted_color

        return slide
