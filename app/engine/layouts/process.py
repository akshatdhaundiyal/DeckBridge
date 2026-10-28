from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from app.engine.coordinates import (
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
    get_column_bounds,
    get_row_bounds,
)
from app.engine.layouts.base import BaseLayoutBuilder
from app.engine.theme import apply_zero_margins_and_wrap, parse_hex_color
from app.schemas.deck import ProcessSlide

SHAPE_TYPE_MAP = {
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "oval": MSO_SHAPE.OVAL,
}


class ProcessLayoutBuilder(BaseLayoutBuilder):
    """Layout builder for vector Process / Flowchart pipelines (horizontal & vertical)."""

    def render(self, slide_model: ProcessSlide):
        slide, _, text_color, muted_color = self.create_blank_slide(slide_model)

        self.render_header(
            slide=slide,
            kicker=slide_model.kicker,
            title=slide_model.title,
            subtitle=slide_model.subtitle,
            text_color=text_color,
            muted_color=muted_color,
        )

        steps = slide_model.steps
        direction = slide_model.direction

        if direction == "vertical":
            self._render_vertical(slide, steps, text_color, muted_color)
        else:
            self._render_horizontal(slide, steps, text_color, muted_color)

        return slide

    def _render_horizontal(self, slide, steps, text_color, muted_color):
        gap = 0.45
        col_lefts, col_w = get_column_bounds(
            total_cols=len(steps),
            left=CONTENT_LEFT,
            total_width=CONTENT_WIDTH,
            gap=gap,
        )

        node_top = CONTENT_TOP + 0.3
        node_h = CONTENT_HEIGHT - 0.6

        for idx, (step, col_left) in enumerate(zip(steps, col_lefts)):
            shape_type = SHAPE_TYPE_MAP.get(step.shape_type, MSO_SHAPE.ROUNDED_RECTANGLE)
            accent = parse_hex_color(step.accent_color) if step.accent_color else self.theme.accent

            # Draw step container node
            node = slide.shapes.add_shape(
                shape_type,
                Inches(col_left),
                Inches(node_top),
                Inches(col_w),
                Inches(node_h),
            )
            if node.has_text_frame:
                apply_zero_margins_and_wrap(node.text_frame)
            node.fill.solid()
            node.fill.fore_color.rgb = self.theme.card_bg
            node.line.color.rgb = self.theme.card_border
            node.line.width = Pt(1.5)

            # Node inner text box
            pad = 0.2
            tb = slide.shapes.add_textbox(
                Inches(col_left + pad),
                Inches(node_top + pad),
                Inches(col_w - pad * 2),
                Inches(node_h - pad * 2),
            )
            tf = tb.text_frame
            apply_zero_margins_and_wrap(tf)

            first_p = True

            # Step Number (e.g. "01")
            step_num = step.step_number if step.step_number is not None else (idx + 1)
            p_num = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_num.text = f"{step_num:02d}"
            p_num.font.name = self.theme.font_family
            p_num.font.size = Pt(14)
            p_num.font.bold = True
            p_num.font.color.rgb = accent
            p_num.space_after = Pt(6)
            first_p = False

            # Step Title
            p_title = tf.paragraphs[0] if first_p else tf.add_paragraph()
            p_title.text = step.title
            p_title.font.name = self.theme.font_family
            p_title.font.size = Pt(15)
            p_title.font.bold = True
            p_title.font.color.rgb = text_color
            p_title.space_after = Pt(6)
            first_p = False

            # Step Description
            if step.body:
                p_body = tf.add_paragraph()
                p_body.text = step.body
                p_body.font.name = self.theme.font_family
                p_body.font.size = Pt(11)
                p_body.font.color.rgb = muted_color

            # Draw Right Arrow to Next Step
            if idx < len(steps) - 1:
                arrow_w = gap * 0.6
                arrow_h = 0.35
                arrow_left = col_left + col_w + (gap - arrow_w) / 2.0
                arrow_top = node_top + (node_h / 2.0) - (arrow_h / 2.0)

                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    Inches(arrow_left),
                    Inches(arrow_top),
                    Inches(arrow_w),
                    Inches(arrow_h),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = self.theme.accent
                arrow.line.fill.background()

    def _render_vertical(self, slide, steps, text_color, muted_color):
        gap = 0.3
        row_tops, row_h = get_row_bounds(
            total_rows=len(steps),
            top=CONTENT_TOP,
            total_height=CONTENT_HEIGHT,
            gap=gap,
        )

        for idx, (step, row_top) in enumerate(zip(steps, row_tops)):
            shape_type = SHAPE_TYPE_MAP.get(step.shape_type, MSO_SHAPE.ROUNDED_RECTANGLE)
            accent = parse_hex_color(step.accent_color) if step.accent_color else self.theme.accent

            node = slide.shapes.add_shape(
                shape_type,
                Inches(CONTENT_LEFT),
                Inches(row_top),
                Inches(CONTENT_WIDTH),
                Inches(row_h),
            )
            if node.has_text_frame:
                apply_zero_margins_and_wrap(node.text_frame)
            node.fill.solid()
            node.fill.fore_color.rgb = self.theme.card_bg
            node.line.color.rgb = self.theme.card_border
            node.line.width = Pt(1.5)

            pad = 0.2
            tb = slide.shapes.add_textbox(
                Inches(CONTENT_LEFT + pad),
                Inches(row_top + pad),
                Inches(CONTENT_WIDTH - pad * 2),
                Inches(row_h - pad * 2),
            )
            tf = tb.text_frame
            apply_zero_margins_and_wrap(tf)

            step_num = step.step_number if step.step_number is not None else (idx + 1)
            p = tf.paragraphs[0]
            p.text = f"{step_num:02d}  •  {step.title}"
            p.font.name = self.theme.font_family
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = text_color

            if step.body:
                p.space_after = Pt(2)
                p_body = tf.add_paragraph()
                p_body.text = step.body
                p_body.font.name = self.theme.font_family
                p_body.font.size = Pt(11)
                p_body.font.color.rgb = muted_color

            # Draw Down Arrow to Next Step
            if idx < len(steps) - 1:
                arrow_w = 0.35
                arrow_h = gap * 0.6
                arrow_left = CONTENT_LEFT + (CONTENT_WIDTH / 2.0) - (arrow_w / 2.0)
                arrow_top = row_top + row_h + (gap - arrow_h) / 2.0

                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.DOWN_ARROW,
                    Inches(arrow_left),
                    Inches(arrow_top),
                    Inches(arrow_w),
                    Inches(arrow_h),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = self.theme.accent
                arrow.line.fill.background()
