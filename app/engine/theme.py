import re
from dataclasses import dataclass
from typing import Optional
from pptx.dml.color import RGBColor
from pptx.util import Inches

from app.core.exceptions import InvalidColorError
from app.schemas.deck import DeckTheme


def parse_hex_color(hex_str: str) -> RGBColor:
    """Sanitize and convert hex string to RGBColor.

    Strips '#' and converts 6-character (or 3-character shorthand) hex strings.
    """
    if not isinstance(hex_str, str):
        raise InvalidColorError(f"Color value must be a string, got {type(hex_str).__name__}")

    clean = hex_str.strip().lstrip("#")

    # Support shorthand 3-char hex like "F00" -> "FF0000"
    if len(clean) == 3 and re.match(r"^[0-9a-fA-F]{3}$", clean):
        clean = "".join(c * 2 for c in clean)

    if len(clean) != 6 or not re.match(r"^[0-9a-fA-F]{6}$", clean):
        raise InvalidColorError(f"Invalid hex color '{hex_str}'. Expected 6-character hex format.")

    r = int(clean[0:2], 16)
    g = int(clean[2:4], 16)
    b = int(clean[4:6], 16)
    return RGBColor(r, g, b)


def apply_zero_margins_and_wrap(text_frame) -> None:
    """Enforce OpenXML compliance: explicit zero-margins and word wrap."""
    text_frame.margin_top = Inches(0)
    text_frame.margin_bottom = Inches(0)
    text_frame.margin_left = Inches(0)
    text_frame.margin_right = Inches(0)
    text_frame.word_wrap = True


@dataclass
class ResolvedTheme:
    primary: RGBColor
    accent: RGBColor
    card_bg: RGBColor
    text: RGBColor
    muted: RGBColor
    background: RGBColor
    card_border: RGBColor
    font_family: str

    @classmethod
    def from_schema(cls, theme: DeckTheme) -> "ResolvedTheme":
        return cls(
            primary=parse_hex_color(theme.primary_color),
            accent=parse_hex_color(theme.accent_color),
            card_bg=parse_hex_color(theme.card_bg),
            text=parse_hex_color(theme.text_color),
            muted=parse_hex_color(theme.muted_color),
            background=parse_hex_color(theme.background_color),
            card_border=parse_hex_color(theme.card_border),
            font_family=theme.font_family,
        )

    def resolve_slide_colors(
        self, background_override: Optional[str]
    ) -> tuple[RGBColor, RGBColor, RGBColor]:
        """Resolve (background_color, text_color, muted_color) with optional slide override."""
        if not background_override:
            return self.background, self.text, self.muted

        override = background_override.strip().lower()
        if override == "dark":
            return parse_hex_color("0F172A"), parse_hex_color("F8FAFC"), parse_hex_color("94A3B8")
        elif override == "light":
            return parse_hex_color("FFFFFF"), parse_hex_color("0F172A"), parse_hex_color("64748B")
        else:
            # Custom hex background
            bg_rgb = parse_hex_color(background_override)
            # Simple luminance check for text contrast: 0.299*R + 0.587*G + 0.114*B
            r, g, b = bg_rgb[0], bg_rgb[1], bg_rgb[2]
            luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
            if luminance < 0.5:
                # Dark background -> light text
                return bg_rgb, parse_hex_color("F8FAFC"), parse_hex_color("94A3B8")
            else:
                return bg_rgb, parse_hex_color("0F172A"), parse_hex_color("64748B")
